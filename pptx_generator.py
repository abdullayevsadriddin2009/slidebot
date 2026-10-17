# -*- coding: utf-8 -*-
"""
PPTX generator — to'liq qayta yozildi.
Namunaviy fayllarga (Amir_Temur.pptx, Hujayra.pptx) 1:1 mos layout:
  - 16:9 (13.333 x 7.5 dyuym)
  - Sarlavha shrifti: Playfair Display (qalin)
  - Matn shrifti: Noto Sans, asosiy matn 15pt (foydalanuvchi talabi)
  - 2 tema: temurid (to'q ko'k #1B2A41 + oltin #C9A227) va biolog (to'q yashil #1B4332 + yashil #4ADE80)
  - Tuzilma: 1 muqova + 1 mundarija + (N-3) ta kontent slayd + 1 yakuniy = N
  - Kontent slaydlarda 6 xil layout (A-F) navbat bilan — bo'sh joy qolmaydi.
  - 6 til: uz / ru / en / tr / de / fr.
Interfeys (o'zgartirilmaydi, creation.py chaqiradi):
  generate_pptx(topic, subject, lang, design_key, slide_count, slides_content, images=None, author="") -> BytesIO
"""
import io
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

try:
    from designs import select_theme, DESIGNS, subject_name, lang_full_name
except Exception:
    select_theme = lambda s, d=None: "temurid"
    DESIGNS = {}
    subject_name = lambda s, l="uz": s
    lang_full_name = lambda c, l="uz": c

BOT_USERNAME = "@Talaba_AI_SlideBot"

# ===== 6 til uchun yorliqlar =====
LABELS = {
    "uz": {"agenda":"MUNDARIJA","conclusion":"XULOSA","thankyou":"E'tiboringiz uchun rahmat!",
           "questions":"Savollar uchun","author_label":"Bajaruvchi","subject_label":"Fan",
           "created_by":"Yaratildi","overview":"UMUMIY KO'RINISH","key_facts":"ASOSIY FAKTLAR",
           "timeline":"JARAYON BOSQICHLARI","comparison":"SOLISHTIRMA","table_note":"Qisqacha jadval",
           "sources":"Manbalar","detail":"Tafsilotlar","key_point":"Asosiy nuqta","quote":"Iqtibos"},
    "ru": {"agenda":"СОДЕРЖАНИЕ","conclusion":"ЗАКЛЮЧЕНИЕ","thankyou":"Спасибо за внимание!",
           "questions":"Для вопросов","author_label":"Выполнил","subject_label":"Предмет",
           "created_by":"Создано","overview":"ОБЗОР","key_facts":"ОСНОВНЫЕ ФАКТЫ",
           "timeline":"ЭТАПЫ","comparison":"СРАВНЕНИЕ","table_note":"Краткая таблица",
           "sources":"Источники","detail":"Детали","key_point":"Главный пункт","quote":"Цитата"},
    "en": {"agenda":"AGENDA","conclusion":"CONCLUSION","thankyou":"Thank you for your attention!",
           "questions":"For questions","author_label":"By","subject_label":"Subject",
           "created_by":"Created by","overview":"OVERVIEW","key_facts":"KEY FACTS",
           "timeline":"STEPS","comparison":"COMPARISON","table_note":"Summary table",
           "sources":"Sources","detail":"Details","key_point":"Key point","quote":"Quote"},
    "tr": {"agenda":"İÇİNDEKİLER","conclusion":"SONUÇ","thankyou":"Dikkatiniz için teşekkürler!",
           "questions":"Sorular için","author_label":"Hazırlayan","subject_label":"Ders",
           "created_by":"Oluşturan","overview":"GENEL BAKIŞ","key_facts":"TEMEL GERÇEKLER",
           "timeline":"AŞAMALAR","comparison":"KARŞILAŞTIRMA","table_note":"Özet tablo",
           "sources":"Kaynaklar","detail":"Ayrıntılar","key_point":"Ana nokta","quote":"Alıntı"},
    "de": {"agenda":"INHALTSVERZEICHNIS","conclusion":"FAZIT","thankyou":"Vielen Dank für Ihre Aufmerksamkeit!",
           "questions":"Für Fragen","author_label":"Von","subject_label":"Fach",
           "created_by":"Erstellt von","overview":"ÜBERSICHT","key_facts":"KERNFakten".replace("Kern","KERN"),
           "timeline":"ABSCHNITTE","comparison":"VERGLEICH","table_note":"Übersichtstabelle",
           "sources":"Quellen","detail":"Details","key_point":"Kernpunkt","quote":"Zitat"},
    "fr": {"agenda":"SOMMAIRE","conclusion":"CONCLUSION","thankyou":"Merci pour votre attention !",
           "questions":"Pour les questions","author_label":"Par","subject_label":"Matière",
           "created_by":"Créé par","overview":"APERÇU","key_facts":"FAITS CLÉS",
           "timeline":"ÉTAPES","comparison":"COMPARAISON","table_note":"Tableau récapitulatif",
           "sources":"Sources","detail":"Détails","key_point":"Point clé","quote":"Citation"},
}

def _lbl(lang):
    return LABELS.get(str(lang).lower(), LABELS["en"])

def _rgb(tup):
    return RGBColor(tup[0], tup[1], tup[2])

def _set_bg(slide, rgb):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(rgb)

def _no_line(shape):
    shape.line.fill.background()

def _rect(slide, x, y, w, h, fill_rgb, line_rgb=None, line_w=0.75, rounded=False):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shp_type, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = _rgb(fill_rgb)
    if line_rgb:
        sp.line.color.rgb = _rgb(line_rgb)
        sp.line.width = Pt(line_w)
    else:
        _no_line(sp)
    try:
        sp.shadow.inherit = False
    except Exception:
        pass
    return sp

def _line(slide, x, y, w, rgb, h=0.04):
    return _rect(slide, x, y, w, h, rgb)

def _txt(slide, x, y, w, h, text, size=15, bold=False, rgb=(40,40,40), font="Noto Sans",
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, line_spacing=1.25):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Emu(0); tf.margin_right = Emu(0); tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = ln
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = _rgb(rgb)
    return tb

def _bullets(slide, x, y, w, h, items, size=15, rgb=(43,43,43), font="Noto Sans",
             bold_labels=True, line_spacing=1.3, bullet_char="•"):
    """Bold label (birinchi so'zgacha ':') + oddiy matn ko'rinishidagi ro'yxat."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0); tf.margin_right = Emu(0); tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        p.space_after = Pt(6)
        clean = str(it).strip().lstrip("•-–—0123456789.) ").strip()
        label, sep, rest = clean.partition(":")
        r0 = p.add_run(); r0.text = f"{bullet_char}  "; r0.font.name=font; r0.font.size=Pt(size); r0.font.bold=True; r0.font.color.rgb=_rgb(rgb)
        if sep and bold_labels and rest.strip():
            r1 = p.add_run(); r1.text = label.strip() + ": "; r1.font.name=font; r1.font.size=Pt(size); r1.font.bold=True; r1.font.color.rgb=_rgb((20,20,20))
            r2 = p.add_run(); r2.text = rest.strip(); r2.font.name=font; r2.font.size=Pt(size); r2.font.bold=False; r2.font.color.rgb=_rgb(rgb)
        else:
            r1 = p.add_run(); r1.text = clean; r1.font.name=font; r1.font.size=Pt(size); r1.font.bold=False; r1.font.color.rgb=_rgb(rgb)
    return tb

def _pic(slide, path, x, y, w, h):
    """Rasm qo'shish — xato bo'lsa None."""
    try:
        if path and os.path.exists(str(path)):
            return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    except Exception:
        pass
    return None

def _decor_image(theme_key, w=1600, h=900):
    """Rasm yo'q bo'lsa — gradient dekor (Pillow orqali). BytesIO qaytaradi."""
    try:
        from PIL import Image, ImageDraw
        if theme_key == "biolog":
            c1, c2 = (27,67,50), (74,222,128)
        else:
            c1, c2 = (10,35,66), (201,162,39)
        img = Image.new("RGB", (w, h), c1)
        dr = ImageDraw.Draw(img)
        for y in range(h):
            t = y / h
            r = int(c1[0]*(1-t) + c2[0]*t)
            g = int(c1[1]*(1-t) + c2[1]*t)
            b = int(c1[2]*(1-t) + c2[2]*t)
            dr.line([(0,y),(w,y)], fill=(r,g,b))
        buf = io.BytesIO(); img.save(buf, format="PNG"); buf.seek(0)
        return buf
    except Exception:
        return None

def _extract_bullets(slide_data):
    if isinstance(slide_data, dict):
        c = slide_data.get("content") or slide_data.get("points") or []
        if isinstance(c, str):
            c = [x.strip() for x in c.split("\n") if x.strip()]
        return list(c)
    if isinstance(slide_data, (list, tuple)):
        return [str(x) for x in slide_data]
    return []

def _pad_bullets(bullets, n, topic, lbl):
    """Yetarli bo'lmasa — mavzudan kelib chiqqan qo'shimcha qatorlar bilan to'ldirish (bo'sh joy qolmasin)."""
    extras = [
        f"{lbl['detail']}: {topic} — nazariy asoslar va amaliy ahamiyati",
        f"{lbl['key_point']}: Tarixiy kontekst va zamonaviy talqin",
        f"{lbl['detail']}: Tahlil qilishning metodologik asoslari",
        f"{lbl['key_point']}: Global tajriba va mahalliy o'xshashliklar",
        f"{lbl['detail']}: Rivojlanish istiqbollari va ta'sir omillari",
        f"{lbl['key_point']}: Xulosa va amaliy tavsiyalar",
    ]
    out = list(bullets)
    i = 0
    while len(out) < n:
        out.append(extras[i % len(extras)])
        i += 1
    return out[:n]

# ===== Umumiy elementlar =====
def _footer(slide, th, page, total, topic, subject, lang):
    sec = (subject_name(subject, lang) or "").upper()
    left = f"{str(topic)[:38].upper()}  ·  {sec}" if sec else str(topic)[:50].upper()
    _txt(slide, 0.69, 7.01, 8.5, 0.3, left, size=11, rgb=(120,120,120), font="Noto Sans")
    _txt(slide, 10.5, 7.01, 2.14, 0.3, f"{page:02d} / {total:02d}", size=11, rgb=(120,120,120),
         font="Noto Sans", align=PP_ALIGN.RIGHT)

def _content_header(slide, th, eyebrow, title, accent_bright):
    _rect(slide, 0, 0, 0.08, 7.5, accent_bright)
    _txt(slide, 0.69, 0.49, 12.0, 0.35, eyebrow, size=12, bold=True, rgb=accent_bright, font="Noto Sans")
    _txt(slide, 0.69, 0.85, 12.0, 0.75, title, size=28, bold=True, rgb=th["primary"], font=th["title_font"])
    _line(slide, 0.69, 1.68, 0.83, accent_bright, h=0.04)

# ===== 1. Muqova slayd =====
def _cover(prs, topic, subject, lang, th, accent_bright, images, author):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, th["dark_close"])
    cover_img = images[0] if (isinstance(images, list) and images) else images
    pic = _pic(slide, cover_img, 0, 0, 13.333, 7.5)
    if not pic:
        buf = _decor_image(th.get("style","temurid"))
        if buf:
            try:
                slide.shapes.add_picture(buf, 0, 0, Inches(13.333), Inches(7.5))
            except Exception:
                pass
    ov = _rect(slide, 0, 0, 13.333, 7.5, th["dark_close"])
    lbl = _lbl(lang)
    subj = subject_name(subject, lang)
    _rect(slide, 1.11, 2.0, 0.11, 3.2, accent_bright)
    _txt(slide, 1.45, 2.08, 10.5, 0.4, f"{subj.upper()}  ·  {lang_full_name(lang, lang).upper()}",
         size=16, bold=True, rgb=accent_bright, font="Noto Sans")
    words = str(topic).strip().split()
    if len(words) > 1:
        main = " ".join(words[:-1]); last = words[-1]
        tb = slide.shapes.add_textbox(Inches(1.45), Inches(2.6), Inches(11.0), Inches(2.2))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.line_spacing = 1.05
        r1 = p.add_run(); r1.text = main + " "; r1.font.name=th["title_font"]; r1.font.size=Pt(54); r1.font.bold=True; r1.font.color.rgb=_rgb((255,255,255))
        r2 = p.add_run(); r2.text = last; r2.font.name=th["title_font"]; r2.font.size=Pt(54); r2.font.bold=True; r2.font.color.rgb=_rgb(accent_bright)
    else:
        _txt(slide, 1.45, 2.6, 11.0, 2.2, str(topic), size=54, bold=True, rgb=(255,255,255), font=th["title_font"], line_spacing=1.05)
    _line(slide, 1.45, 4.95, 1.67, accent_bright, h=0.05)
    _txt(slide, 1.45, 5.2, 10.5, 0.6, f"{lbl['overview']} · {subj}", size=18, rgb=(240,240,240), font="Noto Sans")
    if author and str(author) not in ("-",""):
        _txt(slide, 1.45, 5.85, 10.5, 0.4, f"{lbl['author_label']}: {author}", size=14, rgb=(220,220,220), font="Noto Sans")
    _txt(slide, 1.45, 6.81, 10.5, 0.3, f"{lbl['created_by']} {BOT_USERNAME}", size=12, rgb=(180,180,180), font="Noto Sans")

# ===== 2. Mundarija slayd =====
def _agenda(prs, slides_content, lang, th, accent_bright, total, topic, subject):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, th["bg"])
    _rect(slide, 0, 0, 13.333, 0.08, accent_bright)
    lbl = _lbl(lang)
    subj = subject_name(subject, lang)
    _txt(slide, 0.83, 0.56, 11.67, 0.35, f"{subj.upper()}  ·  {str(topic)[:40].upper()}", size=13, bold=True, rgb=accent_bright, font="Noto Sans")
    _txt(slide, 0.83, 1.0, 11.0, 0.7, lbl["agenda"], size=34, bold=True, rgb=th["primary"], font=th["title_font"])
    _line(slide, 0.83, 1.9, 0.83, accent_bright, h=0.04)
    _rect(slide, 0.83, 2.4, 11.67, 3.9, (255,255,255), line_rgb=accent_bright, line_w=1.2, rounded=True)
    _rect(slide, 0.83, 2.4, 11.67, 0.62, accent_bright)
    _txt(slide, 1.1, 2.52, 11.0, 0.4, lbl["overview"].upper(), size=18, bold=True, rgb=(255,255,255), font="Noto Sans")
    titles = []
    for sd in (slides_content or []):
        t = (sd.get("title") if isinstance(sd, dict) else None) or ""
        if t and t not in titles:
            titles.append(t)
    if len(titles) < 6:
        titles = _pad_bullets(titles, 6, topic, lbl)
    n_items = min(len(titles), 8)
    half = (n_items + 1) // 2
    col_w = 5.3
    for col in range(2):
        items = titles[col*half:(col+1)*half]
        x = 1.1 + col * (col_w + 0.4)
        tb = slide.shapes.add_textbox(Inches(x), Inches(3.25), Inches(col_w), Inches(2.9))
        tf = tb.text_frame; tf.word_wrap = True
        for i, t in enumerate(items):
            p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.line_spacing = 1.2; p.space_after = Pt(8)
            num = col*half + i + 1
            r1 = p.add_run(); r1.text = f"{num:02d}.  "; r1.font.name=th["title_font"]; r1.font.size=Pt(16); r1.font.bold=True; r1.font.color.rgb=_rgb(accent_bright)
            r2 = p.add_run(); r2.text = str(t)[:55]; r2.font.name="Noto Sans"; r2.font.size=Pt(15); r2.font.color.rgb=_rgb((43,43,43))
    _footer(slide, th, 2, total, topic, subject, lang)

# ===== Kontent layoutlari =====
def _layout_A(slide, th, accent_bright, title, bullets, images, img_idx, lbl, topic):
    b = _pad_bullets(bullets, 5, topic, lbl)
    img_path = images[img_idx % len(images)] if images else None
    if not _pic(slide, img_path, 0.69, 2.01, 3.33, 4.58):
        _rect(slide, 0.69, 2.01, 3.33, 4.58, th["primary"])
        _txt(slide, 0.85, 3.8, 3.0, 1.0, lbl["detail"].upper(), size=16, bold=True, rgb=accent_bright, font="Noto Sans", align=PP_ALIGN.CENTER)
    _txt(slide, 0.69, 6.62, 3.33, 0.3, f"{lbl['sources']}: {BOT_USERNAME}", size=11, rgb=(120,120,120), font="Noto Sans")
    _txt(slide, 4.44, 2.01, 8.19, 0.45, str(b[0])[:70], size=17, bold=True, rgb=th["primary"], font="Noto Sans")
    _line(slide, 4.44, 2.52, 8.19, (220,220,220), h=0.02)
    _bullets(slide, 4.44, 2.7, 8.19, 3.3, b[1:5], size=15, rgb=(43,43,43))
    _rect(slide, 4.44, 6.15, 8.19, 0.62, th["bg"], line_rgb=accent_bright, line_w=1.0)
    _txt(slide, 4.64, 6.28, 7.8, 0.4, f"“{str(b[-1])[:90]}”", size=12, italic=True, rgb=th["primary"], font="Noto Sans")

def _layout_B(slide, th, accent_bright, title, bullets, images, img_idx, lbl, topic):
    b = _pad_bullets(bullets, 6, topic, lbl)
    img_path = images[img_idx % len(images)] if images else None
    if not _pic(slide, img_path, 0.69, 2.01, 7.22, 4.17):
        _rect(slide, 0.69, 2.01, 7.22, 4.17, th["primary"])
        _txt(slide, 0.85, 3.7, 6.8, 1.0, lbl["overview"].upper(), size=18, bold=True, rgb=accent_bright, font="Noto Sans", align=PP_ALIGN.CENTER)
    _txt(slide, 0.69, 6.22, 7.22, 0.3, f"{lbl['detail']} — {str(title)[:50]}", size=11, rgb=(120,120,120), font="Noto Sans")
    _rect(slide, 8.33, 2.01, 4.31, 4.17, (255,255,255), line_rgb=accent_bright, line_w=1.0, rounded=True)
    _rect(slide, 8.33, 2.01, 4.31, 0.5, accent_bright)
    _txt(slide, 8.5, 2.1, 4.0, 0.35, lbl["key_facts"], size=14, bold=True, rgb=(255,255,255), font="Noto Sans")
    _bullets(slide, 8.55, 2.65, 3.9, 3.4, b[:5], size=13, rgb=(43,43,43), bullet_char="▸")

def _layout_C(slide, th, accent_bright, title, bullets, lbl, topic):
    b = _pad_bullets(bullets, 8, topic, lbl)
    _rect(slide, 0.69, 2.01, 5.56, 4.58, (255,255,255), line_rgb=accent_bright, line_w=1.0, rounded=True)
    _rect(slide, 0.69, 2.01, 5.56, 0.55, accent_bright)
    _txt(slide, 0.9, 2.12, 5.2, 0.35, lbl["key_point"].upper(), size=15, bold=True, rgb=(255,255,255), font="Noto Sans")
    _txt(slide, 0.95, 2.75, 5.1, 3.7, str(b[0])[:400], size=15, rgb=(43,43,43), font="Noto Sans", line_spacing=1.4)
    tb = slide.shapes.add_textbox(Inches(6.81), Inches(2.01), Inches(5.83), Inches(4.58))
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(b[1:7]):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.line_spacing = 1.25; p.space_after = Pt(10)
        r1 = p.add_run(); r1.text = f"{i+1:02d}   "; r1.font.name=th["title_font"]; r1.font.size=Pt(18); r1.font.bold=True; r1.font.color.rgb=_rgb(accent_bright)
        r2 = p.add_run(); r2.text = str(it)[:90]; r2.font.name="Noto Sans"; r2.font.size=Pt(15); r2.font.color.rgb=_rgb((43,43,43))

def _layout_D(slide, th, accent_bright, title, bullets, lbl, topic):
    b = _pad_bullets(bullets, 6, topic, lbl)
    nums = ["01","02","03","04","05","06"]
    xs = [0.69, 4.79, 8.89]; ys = [2.01, 4.31]
    idx = 0
    for y in ys:
        for x in xs:
            _rect(slide, x, y, 3.75, 2.08, (255,255,255), line_rgb=accent_bright, line_w=1.0, rounded=True)
            col = accent_bright if idx % 2 == 0 else th["primary"]
            _txt(slide, x+0.25, y+0.2, 3.25, 0.7, nums[idx], size=32, bold=True, rgb=col, font=th["title_font"])
            _line(slide, x+0.25, y+0.95, 0.6, accent_bright, h=0.03)
            _txt(slide, x+0.25, y+1.08, 3.25, 0.9, str(b[idx])[:80], size=12, rgb=(60,60,60), font="Noto Sans", line_spacing=1.2)
            idx += 1

def _layout_E(slide, th, accent_bright, title, bullets, lbl, topic):
    b = _pad_bullets(bullets, 6, topic, lbl)
    xs = [0.69, 3.78, 6.87, 9.96]
    step_word = lbl["timeline"].split()[0][:8]
    for i, x in enumerate(xs):
        _rect(slide, x, 2.01, 2.78, 1.94, (255,255,255), line_rgb=accent_bright, line_w=1.0, rounded=True)
        _rect(slide, x, 2.01, 2.78, 0.49, accent_bright)
        _txt(slide, x+0.15, 2.1, 2.5, 0.35, f"{step_word} {i+1}", size=14, bold=True, rgb=(255,255,255), font="Noto Sans", align=PP_ALIGN.CENTER)
        _txt(slide, x+0.18, 2.62, 2.45, 1.25, str(b[i])[:110], size=12, rgb=(50,50,50), font="Noto Sans", line_spacing=1.2)
    _rect(slide, 0.69, 4.31, 11.94, 2.22, th["bg"], line_rgb=accent_bright, line_w=1.2, rounded=True)
    _txt(slide, 0.95, 4.45, 11.4, 0.4, lbl["comparison"].upper(), size=16, bold=True, rgb=th["primary"], font="Noto Sans")
    _line(slide, 0.95, 4.9, 1.0, accent_bright, h=0.03)
    _bullets(slide, 0.95, 5.05, 5.6, 1.4, b[4:5], size=14, rgb=(43,43,43), bullet_char="▸")
    _bullets(slide, 6.8, 5.05, 5.6, 1.4, b[5:6], size=14, rgb=(43,43,43), bullet_char="▸")

def _layout_F(slide, th, accent_bright, title, bullets, lbl, topic):
    b = _pad_bullets(bullets, 5, topic, lbl)
    rows, cols = 5, 2
    tbl_shape = slide.shapes.add_table(rows, cols, Inches(0.69), Inches(2.01), Inches(11.94), Inches(4.3))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(3.5); tbl.columns[1].width = Inches(8.44)
    headers = [lbl["key_point"], lbl["detail"]]
    for c in range(cols):
        cell = tbl.cell(0, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = _rgb(th["primary"])
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; r = p.add_run(); r.text = headers[c]
        r.font.name="Noto Sans"; r.font.size=Pt(14); r.font.bold=True; r.font.color.rgb=_rgb((255,255,255))
    for r_i in range(1, rows):
        item = str(b[r_i-1]) if r_i-1 < len(b) else ""
        label, sep, rest = item.partition(":")
        vals = [label.strip()[:40] if sep else lbl["key_point"]+f" {r_i}", (rest.strip() if sep else item)[:120]]
        for c in range(cols):
            cell = tbl.cell(r_i, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb((245,247,250) if r_i % 2 else (255,255,255))
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; r = p.add_run(); r.text = vals[c]
            r.font.name="Noto Sans"; r.font.size=Pt(13); r.font.bold=(c==0); r.font.color.rgb=_rgb((40,40,40))
    _txt(slide, 0.69, 6.4, 11.94, 0.3, f"{lbl['table_note']} — {str(title)[:50]}", size=11, rgb=(120,120,120), font="Noto Sans", italic=True)

# ===== Yakuniy slayd =====
def _closing(prs, topic, subject, lang, th, accent_bright, slides_content, author, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, th["dark_close"])
    lbl = _lbl(lang)
    _txt(slide, 1.11, 1.11, 11.0, 0.4, f"{subject_name(subject, lang).upper()}  ·  {str(topic)[:40].upper()}",
         size=14, bold=True, rgb=accent_bright, font="Noto Sans")
    _txt(slide, 1.11, 1.75, 11.0, 0.9, lbl["conclusion"], size=44, bold=True, rgb=(255,255,255), font=th["title_font"])
    _line(slide, 1.11, 3.05, 1.39, accent_bright, h=0.05)
    _rect(slide, 1.11, 3.4, 11.11, 2.5, (255,255,255), line_rgb=accent_bright, line_w=1.2, rounded=True)
    concl = []
    if slides_content:
        last = slides_content[-1]
        if isinstance(last, dict):
            concl = _extract_bullets(last)[:4]
    concl = _pad_bullets(concl, 4, topic, lbl)
    tb = slide.shapes.add_textbox(Inches(1.4), Inches(3.6), Inches(10.6), Inches(2.1))
    tf = tb.text_frame; tf.word_wrap = True
    for i, c in enumerate(concl):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.line_spacing = 1.3; p.space_after = Pt(8)
        r1 = p.add_run(); r1.text = "✓  "; r1.font.name="Noto Sans"; r1.font.size=Pt(15); r1.font.bold=True; r1.font.color.rgb=_rgb(accent_bright)
        r2 = p.add_run(); r2.text = str(c)[:130]; r2.font.name="Noto Sans"; r2.font.size=Pt(15); r2.font.color.rgb=_rgb((40,40,40))
    _txt(slide, 1.11, 6.15, 11.11, 0.6, lbl["thankyou"], size=28, bold=True, rgb=(255,255,255), font=th["title_font"], align=PP_ALIGN.CENTER)
    _txt(slide, 1.11, 6.95, 11.11, 0.3, f"{lbl['created_by']} {BOT_USERNAME}", size=11, rgb=th["light_text"], font="Noto Sans", align=PP_ALIGN.CENTER)

# ===== ASOSIY FUNKSIYA (interfeys saqlanadi) =====
def generate_pptx(topic, subject, lang, design_key, slide_count, slides_content, images=None, author=""):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    theme_key = select_theme(subject, design_key)
    th = DESIGNS.get(theme_key) or DESIGNS.get("temurid") or {
        "bg":(255,250,240),"primary":(27,42,65),"secondary":(120,120,120),"accent":(201,162,39),
        "title_font":"Playfair Display","body_font":"Noto Sans","style":"temurid",
        "dark_close":(27,42,65),"light_text":(220,225,230),
    }
    accent_bright = th.get("accent_bright", th["accent"])
    lbl = _lbl(lang)

    total = int(slide_count) if slide_count else 10
    if total < 5: total = 10
    if total > 30: total = 25

    img_list = images if isinstance(images, list) else ([images] if images else [])
    img_idx = 1

    _cover(prs, topic, subject, lang, th, accent_bright, img_list, author)
    _agenda(prs, slides_content, lang, th, accent_bright, total, topic, subject)

    content_count = total - 3
    layouts = ["A","B","C","D","E","F"]
    sc = slides_content or []
    for i in range(content_count):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _set_bg(slide, th["bg"])
        sd = sc[i] if i < len(sc) else (sc[-1] if sc else {"title": topic, "content": []})
        stitle = (sd.get("title") if isinstance(sd, dict) else None) or f"{lbl['detail']} {i+1}"
        bullets = _extract_bullets(sd)
        eyebrow = f"{subject_name(subject, lang).upper()}  ·  {lbl['detail'].upper()} {i+3:02d}"
        _content_header(slide, th, eyebrow, str(stitle)[:70], accent_bright)
        L = layouts[i % len(layouts)]
        if L == "A":
            _layout_A(slide, th, accent_bright, stitle, bullets, img_list, img_idx, lbl, topic); img_idx += 1
        elif L == "B":
            _layout_B(slide, th, accent_bright, stitle, bullets, img_list, img_idx, lbl, topic); img_idx += 1
        elif L == "C":
            _layout_C(slide, th, accent_bright, stitle, bullets, lbl, topic)
        elif L == "D":
            _layout_D(slide, th, accent_bright, stitle, bullets, lbl, topic)
        elif L == "E":
            _layout_E(slide, th, accent_bright, stitle, bullets, lbl, topic)
        else:
            _layout_F(slide, th, accent_bright, stitle, bullets, lbl, topic)
        _footer(slide, th, i + 3, total, topic, subject, lang)

    _closing(prs, topic, subject, lang, th, accent_bright, sc, author, total)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
